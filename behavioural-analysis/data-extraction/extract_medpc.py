
import argparse
import os
import re
from openpyxl import Workbook

# -----------------------------
# INPUT AND OUTPUT PATHS
# -----------------------------
parser = argparse.ArgumentParser(
    description="Extract MED-PC event data into individual and combined Excel files."
)
parser.add_argument("input_folder", help="Folder containing MED-PC .txt files.")
parser.add_argument(
    "--output-folder",
    help="Destination folder (default: the input folder).",
)
args = parser.parse_args()

input_folder = os.path.abspath(os.path.expanduser(args.input_folder))
output_folder = os.path.abspath(os.path.expanduser(args.output_folder or input_folder))
if not os.path.isdir(input_folder):
    raise NotADirectoryError(f"Input folder not found: {input_folder}")
os.makedirs(output_folder, exist_ok=True)
combined_output_file = os.path.join(output_folder, "Combined_MEDPC_Processed.xlsx")

# -----------------------------
# HELPER FUNCTIONS
# -----------------------------
def clean_array(raw_text):
    """Remove index numbers and split values into a list of floats."""
    cleaned = re.sub(r"\d+:", "", raw_text)  # Remove indices like '0:'
    return [float(v) for v in cleaned.split() if v.strip()]

def process_medpc_file(file_path):
    """Process a single MED-PC file and return metadata and event timestamps."""
    with open(file_path, "r") as f:
        content = f.read()

    # Extract metadata
    metadata_keys = ["Subject", "Experiment", "Start Date", "End Date", "Start Time", "End Time", "MSN", "A", "C", "M"]
    metadata = {}
    for key in metadata_keys:
        pattern = rf"{key}:(.*)"
        match = re.search(pattern, content)
        if match:
            metadata[key] = match.group(1).strip()

    # Extract arrays Y and Z
    y_match = re.search(r"Y:(.*?)(?=Z:)", content, re.S)
    z_match = re.search(r"Z:(.*)", content, re.S)

    if not y_match or not z_match:
        print(f"⚠️ Warning: Could not find arrays Y and Z in {file_path}")
        return metadata, [], [], []

    y_values = [int(v) for v in clean_array(y_match.group(1))]
    z_values = clean_array(z_match.group(1))

    # Remove trailing zeros from Z to match Y length
    while len(z_values) > len(y_values):
        z_values.pop()

    if len(y_values) != len(z_values):
        print(f"⚠️ Warning: Arrays Y and Z have different lengths in {file_path}")
        return metadata, [], [], []

    # Separate events by code
    left_presses = []
    right_presses = []
    magazine_checks = []

    for code, timestamp in zip(y_values, z_values):
        if code == 3:
            left_presses.append(timestamp)
        elif code == 4:
            right_presses.append(timestamp)
        elif code == 7:
            magazine_checks.append(timestamp)

    return metadata, left_presses, right_presses, magazine_checks

def save_to_excel(metadata, left_presses, right_presses, magazine_checks, output_file):
    """Save metadata and event data to a single-sheet Excel file."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"

    # Write metadata
    ws.append(["Metadata"])
    for k, v in metadata.items():
        ws.append([k, v])
    ws.append([])  # Blank line

    # Write headers for event data
    ws.append([
        "#", "Left Lever Press (s)", "", "#", "Right Lever Press (s)", "", "#", "Magazine Check (s)"
    ])

    # Write event data in parallel columns
    max_len = max(len(left_presses), len(right_presses), len(magazine_checks))
    for i in range(max_len):
        row = [
            i + 1 if i < len(left_presses) else "",
            left_presses[i] if i < len(left_presses) else "",
            "",
            i + 1 if i < len(right_presses) else "",
            right_presses[i] if i < len(right_presses) else "",
            "",
            i + 1 if i < len(magazine_checks) else "",
            magazine_checks[i] if i < len(magazine_checks) else ""
        ]
        ws.append(row)

    wb.save(output_file)
    print(f"✅ Saved: {output_file}")

# -----------------------------
# MAIN PROCESSING LOOP
# -----------------------------
combined_wb = Workbook()
combined_wb.remove(combined_wb.active)  # Remove default sheet

for filename in os.listdir(input_folder):
    if filename.lower().endswith(".txt"):
        file_path = os.path.join(input_folder, filename)
        metadata, left_presses, right_presses, magazine_checks = process_medpc_file(file_path)

        # Save individual file
        base_name = os.path.splitext(filename)[0]
        output_file = os.path.join(output_folder, f"{base_name}_processed.xlsx")
        save_to_excel(metadata, left_presses, right_presses, magazine_checks, output_file)

        # Add to combined workbook
        ws = combined_wb.create_sheet(title=metadata.get("Subject", base_name)[:31])
        ws.append(["Metadata"])
        for k, v in metadata.items():
            ws.append([k, v])
        ws.append([])
        ws.append([
            "#", "Left Lever Press (s)", "", "#", "Right Lever Press (s)", "", "#", "Magazine Check (s)"
        ])
        max_len = max(len(left_presses), len(right_presses), len(magazine_checks))
        for i in range(max_len):
            row = [
                i + 1 if i < len(left_presses) else "",
                left_presses[i] if i < len(left_presses) else "",
                "",
                i + 1 if i < len(right_presses) else "",
                right_presses[i] if i < len(right_presses) else "",
                "",
                i + 1 if i < len(magazine_checks) else "",
                magazine_checks[i] if i < len(magazine_checks) else ""
            ]
            ws.append(row)

combined_wb.save(combined_output_file)
print(f"✅ Combined file saved: {combined_output_file}")
print("✅ All files processed successfully!")
